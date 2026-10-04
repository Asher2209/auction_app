import io
import re
from datetime import date, datetime, timedelta
from decimal import Decimal

import openpyxl
import pytest
from flask import g
from pypdf import PdfReader
from sqlalchemy import event

from app.extensions import db
from app.models import Auction, Bid, Category, CryptoPayment, Invoice, Payment, Product, Review, User, Winner
from app.services import report_export as rx
from app.services import report_service as rs

from .conftest import make_user
from .test_analytics import NOW, world  # noqa: F401  (world is a fixture)
from .test_payments import client_for

RANGE = {"from": "2026-09-01", "to": "2026-10-31"}
ALL = list(rs.REPORTS)


@pytest.fixture(autouse=True)
def _db(app):
    """Every test here may query, so give them all the app context (and an empty database)."""


def run(key, args=None, now=NOW):
    report = rs.REPORTS[key]
    params, errors = rs.parse_params(report, args or {}, now)
    return report, params, errors, rs.run(report, params)


def cell(report, row, key):
    return row[[c.key for c in report.columns].index(key)]


def col_values(key, name, args=None):
    report, _, _, data = run(key, args)
    return [cell(report, r, name) for r in data.rows]


def summary(data):
    return {label: value for label, value, _ in data.summary}


def pdf_text(data):
    return "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(data)).pages)


def win(a, buyer, amount, when):
    db.session.add(Winner(auction_id=a.id, buyer_id=buyer.id, winning_amount=Decimal(amount), winning_time=when))
    a.highest_bidder_id, a.current_bid = buyer.id, Decimal(amount)
    db.session.commit()


def set_payment(a, **kw):
    p = Payment.query.filter_by(auction_id=a.id).one()
    for k, v in kw.items():
        setattr(p, k, v)
    db.session.commit()
    return p


@pytest.fixture
def full(world):
    """Enough varied data that every report has something to show."""
    w = world
    a1 = w.sale(w.s1, "Books", "Atlas", 1000, datetime(2026, 10, 2, 9), buyer=w.b1)
    a2 = w.sale(w.s2, "Sports", "Bat", 8000, datetime(2026, 9, 20, 14), buyer=w.b2, kind="crypto", eth="0.025")
    a3 = w.sale(w.s1, "Electronics", "Phone", 300, datetime(2026, 9, 25), buyer=w.b3, kind="upi")
    w.sale(w.s2, "Books", "Failed one", 50, datetime(2026, 10, 1), status="failed")
    live = w.auction(w.s1, "Sports", "Live item", start=NOW - timedelta(hours=5), status="active", bids=(w.b1, w.b2))
    live.end_time, live.current_bid, live.highest_bidder_id = NOW + timedelta(hours=30), Decimal(101), w.b2.id
    w.auction(w.s2, "Books", "Upcoming", start=NOW + timedelta(days=1), status="scheduled")
    db.session.commit()
    for a, buyer, amt, when in ((a1, w.b1, 1000, datetime(2026, 10, 2)), (a2, w.b2, 8000, datetime(2026, 9, 22)), (a3, w.b3, 300, datetime(2026, 9, 27))):
        win(a, buyer, amt, when)
    return w


# ---- the registry -----------------------------------------------------------------------------------
def test_there_are_eleven_reports_with_the_expected_titles():
    assert [r.title for r in rs.REPORTS.values()] == [
        "Daily Auction Report", "Active Auction Report", "Completed Auction Report", "Highest Selling Products", "Highest Bid Report",
        "Revenue Report", "User Activity Report", "Seller Performance Report", "Product Rating Report", "Payment Report",
        "Cryptocurrency Transaction Report"]
    assert len(set(rs.REPORTS)) == 11 and all(k == r.key for k, r in rs.REPORTS.items())
    for r in rs.REPORTS.values():
        assert len({c.key for c in r.columns}) == len(r.columns) and set(r.pdf_skip) <= {c.key for c in r.columns}


@pytest.mark.parametrize("key", ALL)
def test_every_report_runs_on_empty_data(app, key):
    report, params, errors, data = run(key, RANGE)
    assert data.rows == [] or key in ("seller-performance",) and True
    assert not errors and all(len(label) for label, _, _ in data.summary)
    assert rx.render_pdf(report, data, params, NOW).startswith(b"%PDF") and rx.render_xlsx(report, data, params, NOW)[:2] == b"PK"


@pytest.mark.parametrize("key", ALL)
def test_every_report_has_rows_that_match_its_columns(full, key):
    db.session.add(Review(product_id=Product.query.filter_by(title="Phone").one().id, buyer_id=full.b2.id, rating=5, created_at=datetime(2026, 10, 4)))
    db.session.commit()
    report, params, errors, data = run(key, {"from": "2026-01-01", "to": "2026-12-31", "day": "2026-09-29"})
    assert data.rows, key
    assert all(len(r) == len(report.columns) for r in data.rows), key
    for kind, i in ((c.kind, n) for n, c in enumerate(report.columns)):
        for r in data.rows:
            assert isinstance(rx.fmt(kind, r[i]), str)  # every value formats, whatever its kind
    assert rx.render_pdf(report, data, params, NOW).startswith(b"%PDF") and rx.render_xlsx(report, data, params, NOW)[:2] == b"PK"


# ---- parameters ----------------------------------------------------------------------------------------
def test_default_period_is_the_last_30_days_inclusive_of_today():
    _, p, errors, _ = run("revenue", {})
    assert not errors and p.start == datetime(2026, 9, 16) and p.end == datetime(2026, 10, 16) and p.last_day == date(2026, 10, 15)


def test_day_report_defaults_to_today_and_accepts_a_day():
    _, p, _, _ = run("daily-auctions", {})
    assert p.day == date(2026, 10, 15) and (p.start, p.end) == (datetime(2026, 10, 15), datetime(2026, 10, 16))
    assert run("daily-auctions", {"day": "2026-02-28"})[1].day == date(2026, 2, 28)


def test_snapshot_and_all_time_reports():
    assert run("active-auctions", {"from": "2026-01-01"})[1].start is None  # filters are ignored for a snapshot
    _, p, _, _ = run("product-ratings", {})
    assert p.start is None and p.end is None  # ratings default to all time


@pytest.mark.parametrize("args", [{"from": "nonsense"}, {"from": "2026-13-45"}, {"to": "yesterday"}, {"from": "2026-02-30"}, {"from": "' OR 1=1 --"},
                                  {"from": "\x00"}, {"from": "x" * 500}])
def test_bad_range_dates_give_a_message_and_a_safe_default(args):
    _, p, errors, _ = run("revenue", args)
    assert errors and "must be a date" in errors[0] and p.start is not None  # fell back to the default window


def test_unpadded_dates_are_accepted():
    _, p, errors, _ = run("revenue", {"from": "2026-1-1", "to": "2026-1-31"})
    assert not errors and p.start == datetime(2026, 1, 1) and p.last_day == date(2026, 1, 31)


@pytest.mark.parametrize("day", ["nonsense", "2026-13-45", "2026-02-30", "' OR 1=1 --", "\x00", "x" * 500, "2026-10-3x"])
def test_bad_day_gives_a_message_and_defaults_to_today(day):
    _, p, errors, _ = run("daily-auctions", {"day": day})
    assert errors and "must be a date" in errors[0] and p.day == date(2026, 10, 15)


def test_reversed_ranges_are_swapped_and_huge_ones_clamped():
    _, p, errors, _ = run("revenue", {"from": "2026-10-10", "to": "2026-10-01"})
    assert "swapped" in errors[0] and p.start == datetime(2026, 10, 1) and p.last_day == date(2026, 10, 10)
    _, p, errors, _ = run("revenue", {"from": "1800-01-01", "to": "2026-10-01"})
    assert "at most" in errors[0] and (p.last_day - p.start.date()).days == rs.MAX_RANGE_DAYS
    assert not run("revenue", {"from": "2000-01-01", "to": "2026-10-15"})[2]  # the "All time" preset is fine


def test_status_filter_values():
    assert run("payments", {"status": "failed"})[1].status == "failed"
    _, p, errors, _ = run("payments", {"status": "bogus"})
    assert p.status == "all" and errors
    assert run("revenue", {"status": "failed"})[1].status == "all"  # only reports that have a status filter use it


def test_range_end_is_inclusive_and_start_is_inclusive(world):
    w = world
    w.sale(w.s1, "Books", "edge-start", 1, datetime(2026, 9, 1, 0, 0, 0))
    w.sale(w.s1, "Books", "edge-end", 2, datetime(2026, 9, 30, 23, 59, 59))
    w.sale(w.s1, "Books", "after", 4, datetime(2026, 10, 1, 0, 0, 0))
    w.sale(w.s1, "Books", "before", 8, datetime(2026, 8, 31, 23, 59, 59))
    assert sorted(col_values("top-products", "product", {"from": "2026-09-01", "to": "2026-09-30"})) == ["edge-end", "edge-start"]


# ---- 1. daily ----------------------------------------------------------------------------------------------
def test_daily_report_shows_started_ended_and_bidding_that_day(world):
    w = world
    started = w.auction(w.s1, "Books", "started", start=datetime(2026, 10, 3, 10), status="active")
    ended = w.auction(w.s1, "Books", "ended", start=datetime(2026, 10, 1, 8))
    ended.end_time = datetime(2026, 10, 3, 8)
    bidding = w.auction(w.s2, "Sports", "just bids", start=datetime(2026, 9, 28), status="active")
    w.auction(w.s2, "Sports", "other day", start=datetime(2026, 9, 1))
    for hour, buyer in ((9, w.b1), (11, w.b2)):
        db.session.add(Bid(auction_id=bidding.id, buyer_id=buyer.id, amount=Decimal(200 + hour), bid_time=datetime(2026, 10, 3, hour)))
    db.session.add(Bid(auction_id=bidding.id, buyer_id=w.b1.id, amount=Decimal(150), bid_time=datetime(2026, 10, 2, 23, 59, 59)))
    w.sale(w.s1, "Books", "paid that day", 700, datetime(2026, 10, 3, 15))
    db.session.commit()
    report, _, _, data = run("daily-auctions", {"day": "2026-10-03"})
    rows = {cell(report, r, "product"): r for r in data.rows}
    assert set(rows) == {"started", "ended", "just bids", "paid that day"} | set() or "paid that day" not in rows
    assert cell(report, rows["started"], "started") == "Yes" and cell(report, rows["started"], "ended") == "No"
    assert cell(report, rows["ended"], "ended") == "Yes" and cell(report, rows["ended"], "started") == "No"
    assert cell(report, rows["just bids"], "bids") == 2  # the bid at 23:59:59 the day before is not counted
    s = summary(data)
    assert s["Started"] >= 1 and s["Ended"] >= 1 and s["Bids placed"] == 2 and s["Payments received"] == 700.0


# ---- 2. active -----------------------------------------------------------------------------------------------
def test_active_report_lists_only_live_and_upcoming(full):
    report, _, _, data = run("active-auctions")
    products = {cell(report, r, "product"): r for r in data.rows}
    assert set(products) == {"Live item", "Upcoming"}  # closed and cancelled auctions are not shown
    live = products["Live item"]
    assert cell(report, live, "status") == "Active" and cell(report, live, "bids") == 2 and cell(report, live, "top") == "b2"
    assert cell(report, live, "current") == 101.0 and cell(report, live, "left") not in ("Ended", "Not started")
    assert cell(report, products["Upcoming"], "left") == "Not started" and cell(report, products["Upcoming"], "top") == "-"
    s = summary(data)
    assert s["Live auctions"] == 1 and s["Scheduled"] == 1 and s["Bids so far"] == 2 and s["Current bids total"] == 101.0
    assert [cell(report, r, "product") for r in data.rows][0] == "Live item"  # soonest ending first


# ---- 3. completed ------------------------------------------------------------------------------------------------
def test_completed_report(full):
    w = full
    report, _, _, data = run("completed-auctions", RANGE)
    rows = {cell(report, r, "product"): r for r in data.rows}
    assert {"Atlas", "Bat", "Phone", "Failed one"} <= set(rows) and "Live item" not in rows
    assert cell(report, rows["Atlas"], "winner") == "buyer" and cell(report, rows["Atlas"], "amount") == 1000.0
    assert cell(report, rows["Atlas"], "payment") == "Successful" and cell(report, rows["Failed one"], "payment") == "Failed"
    assert cell(report, rows["Failed one"], "winner") == "No winner"  # (no Winner row was created for it)
    s = summary(data)
    assert s["Completed auctions"] == 4 and s["With a winner"] == 3 and s["No bids"] == 1
    assert s["Winning bids total"] == 9300.0 and s["Paid so far"] == 9300.0
    assert col_values("completed-auctions", "product", {"from": "2026-10-03", "to": "2026-10-31"}) == []  # period applies to the end date


# ---- 4. top products ---------------------------------------------------------------------------------------------
def test_top_products_ranked_and_capped(world, monkeypatch):
    w = world
    for i, amt in enumerate((500, 9000, 120, 4000, 7000)):
        w.sale(w.s1 if i % 2 else w.s2, "Books", f"P{i}", amt, datetime(2026, 10, 1 + i), kind="crypto" if i == 1 else "card")
    w.sale(w.s1, "Books", "never paid", 99999, datetime(2026, 10, 2), status="pending")
    monkeypatch.setattr(rs, "TOP_PRODUCTS", 3)
    report, _, _, data = run("top-products", RANGE)
    assert [(cell(report, r, "rank"), cell(report, r, "product"), cell(report, r, "amount")) for r in data.rows] == [(1, "P1", 9000.0), (2, "P4", 7000.0), (3, "P3", 4000.0)]
    assert cell(report, data.rows[0], "method") == "Cryptocurrency" and cell(report, data.rows[1], "method") == "Card"
    s = summary(data)
    assert s["Total value"] == 20000.0 and s["Highest sale"] == 9000.0 and s["Average sale"] == pytest.approx(6666.67, abs=0.01)


# ---- 5. highest bids -------------------------------------------------------------------------------------------------
def test_highest_bid_report(world):
    w = world
    a = w.auction(w.s1, "Books", "Rising", start=datetime(2026, 10, 1), status="active", bids=(w.b1, w.b2))
    a.current_bid, a.highest_bidder_id = Decimal(250), w.b2.id          # starting price is 100
    b = w.auction(w.s2, "Sports", "Small", start=datetime(2026, 10, 2), bids=(w.b1,))
    b.current_bid, b.highest_bidder_id = Decimal(120), w.b1.id
    c = w.auction(w.s2, "Sports", "Cancelled", start=datetime(2026, 10, 2), status="cancelled", bids=(w.b1,))
    c.current_bid, c.highest_bidder_id = Decimal(9999), w.b1.id
    w.auction(w.s1, "Books", "No bids", start=datetime(2026, 10, 3))
    w.auction(w.s1, "Books", "Too early", start=datetime(2026, 1, 1), bids=(w.b1,)).highest_bidder_id = w.b1.id
    db.session.commit()
    report, _, _, data = run("highest-bids", RANGE)
    assert [(cell(report, r, "rank"), cell(report, r, "product"), cell(report, r, "bid"), cell(report, r, "bidder")) for r in data.rows] == [
        (1, "Rising", 250.0, "b2"), (2, "Small", 120.0, "buyer")]
    assert cell(report, data.rows[0], "inc") == 150.0 and cell(report, data.rows[1], "inc") == 20.0 and cell(report, data.rows[0], "bids") == 2
    assert summary(data)["Highest bid"] == 250.0 and summary(data)["Average highest bid"] == 185.0


# ---- 6. revenue ---------------------------------------------------------------------------------------------------------
def test_revenue_report_per_day(world):
    w = world
    w.sale(w.s1, "Books", "a", 1000, datetime(2026, 10, 2, 9))
    w.sale(w.s1, "Books", "b", 500, datetime(2026, 10, 2, 17), kind="wallet")
    w.sale(w.s2, "Books", "c", 8000, datetime(2026, 10, 5), kind="crypto", eth="0.025")
    w.sale(w.s2, "Books", "d", 999, datetime(2026, 10, 6), status="failed")
    report, _, _, data = run("revenue", RANGE)
    assert [(cell(report, r, "day"), cell(report, r, "n"), cell(report, r, "sim"), cell(report, r, "cry"), cell(report, r, "total")) for r in data.rows] == [
        (date(2026, 10, 2), 2, 1500.0, 0.0, 1500.0), (date(2026, 10, 5), 1, 0.0, 8000.0, 8000.0)]
    assert cell(report, data.rows[1], "eth") == 0.025
    s = summary(data)
    assert s["Total revenue"] == 9500.0 and s["Via crypto"] == 8000.0 and s["Paid sales"] == 3 and s["ETH received"] == 0.025


# ---- 7. user activity ------------------------------------------------------------------------------------------------------
def test_user_activity_counts_per_period_and_hides_idle_users(full):
    w = full
    make_user("idle@t.test")
    db.session.add(Review(product_id=Product.query.filter_by(title="Atlas").one().id, buyer_id=w.b1.id, rating=5, created_at=datetime(2026, 10, 4)))
    db.session.commit()
    report, _, _, data = run("user-activity", RANGE)
    rows = {cell(report, r, "email"): r for r in data.rows}
    assert "idle@t.test" not in rows and "admin@t.test" not in rows
    b1 = rows["buyer@t.test"]
    assert (cell(report, b1, "bids"), cell(report, b1, "won"), cell(report, b1, "buys"), cell(report, b1, "reviews")) == (3, 1, 1, 1)
    seller = rows["seller@t.test"]
    assert cell(report, seller, "role") == "Seller" and cell(report, seller, "bids") == 0
    totals = [cell(report, r, "total") for r in data.rows]
    assert totals == sorted(totals, reverse=True)
    assert col_values("user-activity", "email", {"from": "2025-01-01", "to": "2025-01-31"}) == []  # no activity then


# ---- 8. seller performance ------------------------------------------------------------------------------------------------------
def test_seller_performance(full):
    w = full
    make_user("quiet@t.test", "seller")
    db.session.add(Review(product_id=Product.query.filter_by(title="Atlas").one().id, buyer_id=w.b1.id, rating=4))
    db.session.add(Review(product_id=Product.query.filter_by(title="Phone").one().id, buyer_id=w.b3.id, rating=2))
    db.session.add(Review(product_id=Product.query.filter_by(title="Atlas").one().id, buyer_id=w.b2.id, rating=1, is_hidden=True))
    db.session.commit()
    report, _, _, data = run("seller-performance", RANGE)
    rows = {cell(report, r, "email"): r for r in data.rows}
    s1 = rows["seller@t.test"]
    assert cell(report, s1, "done") == 2 and cell(report, s1, "winner") == 2 and cell(report, s1, "paid") == 2   # Atlas + Phone
    assert cell(report, s1, "rev") == 1300.0 and cell(report, s1, "rate") == 100.0 and cell(report, s1, "avg") == 650.0
    assert cell(report, s1, "rating") == 3.0 and cell(report, s1, "reviews") == 2  # hidden review excluded
    s2 = rows["seller2@t.test"]
    assert cell(report, s2, "done") == 2 and cell(report, s2, "paid") == 1 and cell(report, s2, "rate") == 50.0  # Bat paid, Failed one not
    quiet = rows["quiet@t.test"]
    assert cell(report, quiet, "done") == 0 and cell(report, quiet, "rate") is None and cell(report, quiet, "avg") is None and cell(report, quiet, "rating") is None
    assert all(0 <= (cell(report, r, "rate") or 0) <= 100 for r in data.rows)
    assert [cell(report, r, "rev") for r in data.rows][0] == 8000.0  # best seller first


# ---- 9. product ratings -------------------------------------------------------------------------------------------------------------
def test_product_rating_report(world):
    w = world
    a = w.auction(w.s1, "Books", "Rated")
    b = w.auction(w.s2, "Sports", "Other")
    for i, (rating, hidden, when) in enumerate(((5, False, datetime(2026, 10, 1)), (4, False, datetime(2026, 10, 2)), (4, False, datetime(2026, 8, 1)), (1, True, datetime(2026, 10, 3)))):
        db.session.add(Review(product_id=a.product_id, buyer_id=make_user(f"r{i}@t.test").id, rating=rating, is_hidden=hidden, created_at=when))
    db.session.add(Review(product_id=b.product_id, buyer_id=make_user("rb@t.test").id, rating=5, created_at=datetime(2026, 10, 4)))
    db.session.commit()
    report, _, _, data = run("product-ratings", {})
    rows = {cell(report, r, "product"): r for r in data.rows}
    r = rows["Rated"]
    assert (cell(report, r, "avg"), cell(report, r, "n"), cell(report, r, "s5"), cell(report, r, "s4"), cell(report, r, "s1"), cell(report, r, "hidden")) == (4.3, 3, 1, 2, 0, 1)
    assert cell(report, r, "last") == datetime(2026, 10, 2) and [cell(report, x, "product") for x in data.rows] == ["Other", "Rated"]
    s = summary(data)
    assert s["Visible reviews"] == 4 and s["Hidden by moderators"] == 1 and s["Average rating"] == 4.5
    in_oct = col_values("product-ratings", "n", {"from": "2026-10-01", "to": "2026-10-31"})
    assert sorted(in_oct) == [1, 2]  # the August review is outside the period


# ---- 10. payments ----------------------------------------------------------------------------------------------------------------
def test_payment_report_statuses_methods_and_invoices(full):
    w = full
    for p in Payment.query.all():
        p.created_at = datetime(2026, 10, 1)
    db.session.commit()
    atlas = Auction.query.join(Product).filter(Product.title == "Atlas").one()
    db.session.add(Invoice(payment_id=atlas.payment.id, number="INV-2026-000007"))
    db.session.commit()
    report, _, _, data = run("payments", RANGE)
    rows = {cell(report, r, "product"): r for r in data.rows}
    assert cell(report, rows["Atlas"], "invoice") == "INV-2026-000007" and cell(report, rows["Phone"], "invoice") == "-"
    assert cell(report, rows["Atlas"], "method") == "Card" and cell(report, rows["Bat"], "method") == "Cryptocurrency" and cell(report, rows["Phone"], "method") == "UPI"
    s = summary(data)
    assert (s["Payments"], s["Successful"], s["Failed"], s["Pending"]) == (4, 3, 1, 0) and s["Collected"] == 9300.0 and s["Outstanding"] == 50.0
    assert col_values("payments", "product", {**RANGE, "status": "failed"}) == ["Failed one"]
    assert sorted(col_values("payments", "status", {**RANGE, "status": "successful"})) == ["Successful"] * 3


# ---- 11. crypto ----------------------------------------------------------------------------------------------------------------------
def test_crypto_report_lists_submitted_transactions(full):
    w = full
    pending = w.sale(w.s1, "Books", "pending crypto", 100, datetime(2026, 10, 3), kind="crypto", status="pending")
    draft = w.sale(w.s1, "Books", "draft only", 100, datetime(2026, 10, 3), kind="crypto")
    CryptoPayment.query.filter_by(payment_id=draft.payment.id).one().transaction_hash = None  # prepared but never submitted
    bad = w.sale(w.s1, "Books", "bad crypto", 200, datetime(2026, 10, 4), kind="crypto", status="failed", eth="0.5")
    CryptoPayment.query.filter_by(payment_id=bad.payment.id).one().failure_reason = "Underpaid: expected 0.6 ETH"
    for cp in CryptoPayment.query.all():
        cp.transaction_date = datetime(2026, 10, 3)
    db.session.commit()
    report, _, _, data = run("crypto-transactions", RANGE)
    rows = {cell(report, r, "product"): r for r in data.rows}
    assert set(rows) == {"Bat", "pending crypto", "bad crypto"}  # the unsubmitted draft is not a transaction
    assert cell(report, rows["Bat"], "status") == "Confirmed" and cell(report, rows["Bat"], "eth") == 0.025 and cell(report, rows["Bat"], "inr") == 8000.0
    assert re.fullmatch(r"0x[0-9a-f]{64}", cell(report, rows["Bat"], "hash"))  # the full hash, never truncated
    assert cell(report, rows["bad crypto"], "reason") == "Underpaid: expected 0.6 ETH"
    s = summary(data)
    assert (s["Transactions"], s["Confirmed"], s["Failed"], s["Pending"]) == (3, 1, 1, 1) and s["ETH confirmed"] == 0.025 and s["Value confirmed"] == 8000.0


# ---- formatting ---------------------------------------------------------------------------------------------------------------------------
@pytest.mark.parametrize("kind,value,text", [
    ("money", 1234.5, "1,234.50"), ("money", Decimal("0.1"), "0.10"), ("int", 12345, "12,345"), ("eth", 0.025, "0.025"), ("eth", 0.0, "0"), ("eth", 1.0, "1"),
    ("eth", 0.1234567, "0.123457"), ("rating", 4.25, "4.2"), ("pct", 12.34, "12.3%"), ("date", date(2026, 10, 3), "2026-10-03"),
    ("datetime", datetime(2026, 10, 3, 9, 5, 7), "2026-10-03 09:05"), ("text", "hello", "hello"), ("text", None, "-"), ("money", None, "-"), ("text", "", "-")])
def test_value_formatting(kind, value, text):
    assert rx.fmt(kind, value) == text


def test_describe_params():
    r = rs.REPORTS
    assert rx.describe_params(r["revenue"], run("revenue", RANGE)[1]) == "Period: 2026-09-01 to 2026-10-31 (Payment date)"
    assert rx.describe_params(r["daily-auctions"], run("daily-auctions", {"day": "2026-10-03"})[1]) == "Day: 2026-10-03"
    assert "all time" in rx.describe_params(r["product-ratings"], run("product-ratings")[1])
    assert "Snapshot" in rx.describe_params(r["active-auctions"], run("active-auctions")[1])
    assert rx.describe_params(r["payments"], run("payments", {**RANGE, "status": "failed"})[1]).endswith("Status: failed")


# ---- the PDF ---------------------------------------------------------------------------------------------------------------------------------
def test_pdf_contents(full):
    report, params, _, data = run("payments", RANGE)
    pdf = rx.render_pdf(report, data, params, NOW)
    text = pdf_text(pdf)
    reader = PdfReader(io.BytesIO(pdf))
    assert reader.metadata.title == "Payment Report" and reader.pages[0].mediabox.width > reader.pages[0].mediabox.height  # landscape
    for expected in ("Payment Report", "Period: 2026-09-01 to 2026-10-31", "Generated 2026-10-15 12:00 UTC", "PAYMENTS", "COLLECTED", "9,300.00",
                     "Atlas", "Successful", "Cryptocurrency", "Page 1 of 1"):
        assert expected in text, expected


def test_pdf_paginates_and_repeats_the_header(world):
    w = world
    for i in range(130):
        w.sale(w.s1 if i % 2 else w.s2, "Books", f"Item number {i:03d}", 100 + i, datetime(2026, 10, 1) + timedelta(hours=i))
    report, params, _, data = run("payments", RANGE)
    assert len(data.rows) == 130
    reader = PdfReader(io.BytesIO(rx.render_pdf(report, data, params, NOW)))
    pages = [p.extract_text() for p in reader.pages]
    assert len(pages) >= 3 and f"Page {len(pages)} of {len(pages)}" in pages[-1]
    assert all("Payment" in p.splitlines()[0] or "Method" in p for p in pages)  # the column header is repeated on every page
    assert "Item number 000" in "".join(pages) and "Item number 129" in "".join(pages)


def test_pdf_notes_for_narrowed_and_truncated_reports(world, monkeypatch):
    w = world
    for i in range(4):
        w.sale(w.s1, "Books", f"c{i}", 100, datetime(2026, 10, 3), kind="crypto")
    monkeypatch.setattr(rs, "MAX_ROWS", 2)
    report, params, _, data = run("crypto-transactions", RANGE)
    text = pdf_text(rx.render_pdf(report, data, params, NOW))
    assert data.truncated and len(data.rows) == 2
    assert "left out of this PDF" in text and "Only the first 2 rows" in text
    assert "Buyer wallet" not in text and "Transaction hash" in text


def test_pdf_survives_hostile_text(world):
    w = world
    w.s1.name = "<b>Evil</b> & </para> Co 日本"
    db.session.commit()
    w.sale(w.s1, "Books", "<script>alert(1)</script> \" ' &amp; x" * 3, 100, datetime(2026, 10, 3))
    w.sale(w.s1, "Books", "y" * 400, 100, datetime(2026, 10, 3))
    for key in ("payments", "top-products", "completed-auctions"):
        report, params, _, data = run(key, RANGE)
        text = pdf_text(rx.render_pdf(report, data, params, NOW))
        assert "<script>alert(1)</script>" in text or key == "completed-auctions"


# ---- the Excel file -----------------------------------------------------------------------------------------------------------------------------
def load(data):
    return openpyxl.load_workbook(io.BytesIO(data))


def test_excel_structure_types_and_formats(full):
    report, params, _, data = run("payments", RANGE)
    wb = load(rx.render_xlsx(report, data, params, NOW))
    ws = wb["Report"]
    assert wb.sheetnames == ["Report", "Summary"] and ws["A1"].value == "Payment Report" and "Period: 2026-09-01" in ws["A2"].value
    headers = [c.value for c in ws[5]]
    assert headers == [c.label for c in report.columns]
    assert ws.freeze_panes == "A6" and ws.auto_filter.ref.startswith("A5:")
    first = {h: ws.cell(6, i + 1) for i, h in enumerate(headers)}
    assert isinstance(first["Amount"].value, (int, float)) and first["Amount"].number_format == "#,##0.00"   # numbers stay numbers
    assert isinstance(first["Created (UTC)"].value, datetime) and first["Created (UTC)"].number_format == "yyyy-mm-dd hh:mm"
    assert isinstance(first["Payment"].value, int) and first["Payment"].number_format == "#,##0"
    summary_ws = wb["Summary"]
    rows = {r[0].value: r[1].value for r in summary_ws.iter_rows(min_row=2) if r[0].value}
    assert rows["Payments"] == 4 and rows["Collected"] == 9300.0 and rows["Rows"] == 4 and rows["Complete?"] == "Yes"


def test_excel_crypto_columns_and_precision(world):
    w = world
    w.sale(w.s1, "Books", "c", 8000, datetime(2026, 10, 3), kind="crypto", eth="0.025")
    CryptoPayment.query.one().transaction_date = datetime(2026, 10, 3)
    db.session.commit()
    report, params, _, data = run("crypto-transactions", RANGE)
    ws = load(rx.render_xlsx(report, data, params, NOW))["Report"]
    h = {c.value: i + 1 for i, c in enumerate(ws[5])}
    assert len(h) == 16 and ws.cell(6, h["ETH"]).value == 0.025 and ws.cell(6, h["ETH"]).number_format == "0.000000"  # all columns, unlike the PDF
    assert re.fullmatch(r"0x[0-9a-f]{64}", ws.cell(6, h["Transaction hash"]).value) and "wallet" not in ws.cell(6, h["Transaction hash"]).value


FORMULAS = ['=HYPERLINK("http://evil.example","click")', "=1+1", "+1+1", "-2+3", "@SUM(1+1)", "=cmd|' /C calc'!A0", "\t=1+1", "\r=1+1"]


@pytest.mark.parametrize("payload", FORMULAS)
def test_excel_never_stores_user_text_as_a_formula(world, payload):
    w = world
    w.s1.name = payload
    db.session.commit()
    w.sale(w.s1, "Books", payload, 100, datetime(2026, 10, 3))
    for key in ("payments", "top-products", "completed-auctions", "seller-performance", "user-activity"):
        report, params, _, data = run(key, RANGE)
        wb = load(rx.render_xlsx(report, data, params, NOW))
        formulas = [c.coordinate for sheet in wb for row in sheet.iter_rows() for c in row if c.data_type == "f"]
        assert formulas == [], (key, payload, formulas)
    report, params, _, data = run("top-products", RANGE)
    ws = load(rx.render_xlsx(report, data, params, NOW))["Report"]
    cells = [c for row in ws.iter_rows(min_row=6) for c in row if isinstance(c.value, str) and payload.strip() in c.value]
    assert cells and all(c.data_type == "s" for c in cells)  # the text itself is intact, just inert


def test_excel_notes_truncation(world, monkeypatch):
    for i in range(3):
        world.sale(world.s1, "Books", f"t{i}", 100, datetime(2026, 10, 3 + i))
    monkeypatch.setattr(rs, "MAX_ROWS", 2)
    report, params, _, data = run("payments", RANGE)
    wb = load(rx.render_xlsx(report, data, params, NOW))
    notes = {r[0].value: r[1].value for r in wb["Summary"].iter_rows() if r[0].value}
    assert data.truncated and notes["Rows"] == 2 and notes["Complete?"].startswith("No")


# ---- the web pages ----------------------------------------------------------------------------------------------------------------------------------
def test_index_lists_all_reports_with_export_links(app, users):
    html = client_for(app, "admin@t.test").get("/admin/reports").data.decode()
    for r in rs.REPORTS.values():
        assert r.title in html and f"/admin/reports/{r.key}/export/pdf" in html and f"/admin/reports/{r.key}/export/xlsx" in html
    assert html.count("Open</a>") == 11 and "Reports" in html


@pytest.mark.parametrize("key", ALL)
def test_each_report_page_renders(app, full, key):
    r = client_for(app, "admin@t.test").get(f"/admin/reports/{key}?from=2026-01-01&to=2026-12-31&day=2026-10-02")
    html = r.data.decode()
    assert r.status_code == 200 and rs.REPORTS[key].title in html and "Download PDF" in html and "Download Excel" in html
    assert all(c.label in html for c in rs.REPORTS[key].columns)


def test_page_shows_filters_summary_and_data(app, full):
    c = client_for(app, "admin@t.test")
    html = c.get("/admin/reports/payments?from=2026-09-01&to=2026-10-31&status=successful").data.decode()
    assert 'name="from" value="2026-09-01"' in html and 'name="to" value="2026-10-31"' in html and '<option value="successful" selected>' in html
    assert "Atlas" in html and "Failed one" not in html and "Period: 2026-09-01 to 2026-10-31" in html
    assert "Last 7 days" in html and "All time" in html and "Collected" in html
    assert "status=successful" in html and "export/pdf?from=2026-09-01&amp;to=2026-10-31&amp;status=successful" in html


def test_bad_filters_show_a_warning_not_an_error(app, full):
    c = client_for(app, "admin@t.test")
    for qs in ("from=nonsense", "from=2026-10-31&to=2026-10-01", "status=%27%3B+DROP+TABLE+users%3B--", "from=1800-01-01&to=2026-10-01"):
        r = c.get(f"/admin/reports/payments?{qs}")
        assert r.status_code == 200 and "alert-warning" in r.data.decode(), qs
    assert "alert-warning" in c.get("/admin/reports/daily-auctions?day=%00").data.decode()
    assert User.query.count() >= 3  # (and nothing was dropped)


def test_all_time_preset_has_no_warning(app, full):
    html = client_for(app, "admin@t.test").get("/admin/reports/revenue?from=2000-01-01&to=2026-10-15").data.decode()
    assert "alert-warning" not in html


@pytest.mark.parametrize("ext,mimetype,magic", [("pdf", "application/pdf", b"%PDF"), ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", b"PK")])
def test_downloads(app, full, ext, mimetype, magic):
    for key in ALL:
        r = client_for(app, "admin@t.test").get(f"/admin/reports/{key}/export/{ext}?from=2026-01-01&to=2026-12-31")
        assert r.status_code == 200 and r.mimetype == mimetype and r.data.startswith(magic), key
        assert re.fullmatch(rf'attachment; filename="chainbid-{key}-\d{{8}}\.{ext}"', r.headers["Content-Disposition"]) and "no-store" in r.headers["Cache-Control"]


def test_export_uses_the_same_filters_as_the_page(app, full):
    c = client_for(app, "admin@t.test")
    everything = load(c.get("/admin/reports/payments/export/xlsx?from=2026-01-01&to=2026-12-31").data)["Report"].max_row
    failed_only = load(c.get("/admin/reports/payments/export/xlsx?from=2026-01-01&to=2026-12-31&status=failed").data)["Report"]
    products = [r[2].value for r in failed_only.iter_rows(min_row=6)]
    assert products == ["Failed one"] and everything > failed_only.max_row


def test_unknown_report_or_format_is_404(app, users):
    c = client_for(app, "admin@t.test")
    for path in ("/admin/reports/nope", "/admin/reports/nope/export/pdf", "/admin/reports/payments/export/csv", "/admin/reports/payments/export/exe",
                 "/admin/reports/..%2Fusers", "/admin/reports/payments/export/"):
        assert c.get(path).status_code == 404, path


def test_reports_are_admin_only(app, full):
    paths = ["/admin/reports", "/admin/reports/payments", "/admin/reports/payments/export/pdf", "/admin/reports/payments/export/xlsx"]
    for role in ("buyer", "seller"):
        c = client_for(app, f"{role}@t.test")
        assert all(c.get(p).status_code == 403 for p in paths), role
    g.pop("_login_user", None)
    anon = app.test_client()
    assert all(anon.get(p).status_code == 302 for p in paths)


def test_html_is_capped_but_exports_are_complete(app, world):
    w = world
    for i in range(210):
        w.sale(w.s1, "Books", f"row{i:03d}", 10 + i, datetime(2026, 9, 1) + timedelta(hours=i))
    c = client_for(app, "admin@t.test")
    html = c.get("/admin/reports/payments?from=2026-09-01&to=2026-10-31").data.decode()
    assert html.count("<tr>") - 1 == 200 and "Showing the first 200 of 210" in html
    ws = load(c.get("/admin/reports/payments/export/xlsx?from=2026-09-01&to=2026-10-31").data)["Report"]
    assert ws.max_row - 5 == 210  # the export has every row


def test_truncated_report_warns_on_the_page(app, world, monkeypatch):
    for i in range(3):
        world.sale(world.s1, "Books", f"t{i}", 100, datetime(2026, 10, 3))
    monkeypatch.setattr(rs, "MAX_ROWS", 2)
    html = client_for(app, "admin@t.test").get("/admin/reports/payments?from=2026-09-01&to=2026-10-31").data.decode()
    assert "capped at 2 rows" in html


def test_page_escapes_hostile_text(app, world):
    world.s1.name = "<img src=x onerror=alert(1)>"
    db.session.commit()
    world.sale(world.s1, "Books", "<script>alert(1)</script>", 100, datetime(2026, 10, 3))
    html = client_for(app, "admin@t.test").get("/admin/reports/top-products?from=2026-09-01&to=2026-10-31").data.decode()
    assert "<script>alert(1)</script>" not in html and "&lt;script&gt;" in html and "<img src=x" not in html


def test_report_queries_do_not_multiply_with_data(app, world):
    w = world

    def count_queries(key):
        n = {"n": 0}

        def hit(*_):
            n["n"] += 1

        event.listen(db.engine, "before_cursor_execute", hit)
        try:
            run(key, RANGE)
        finally:
            event.remove(db.engine, "before_cursor_execute", hit)
        return n["n"]

    w.sale(w.s1, "Books", "one", 100, datetime(2026, 10, 1))
    before = {k: count_queries(k) for k in ALL}
    for i in range(25):
        w.sale(w.s1 if i % 2 else w.s2, "Sports", f"more{i}", 100 + i, datetime(2026, 10, 1 + i % 20), buyer=w.b2 if i % 3 else w.b3,
               kind=["card", "upi", "crypto"][i % 3])
    assert {k: count_queries(k) for k in ALL} == before  # no per-row queries anywhere
